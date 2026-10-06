"""Export the supplied item data and check them against the saved JASP project.

This script does not fit CFA models or reproduce the original HTMT analysis.
It preserves original files and writes derived CSV/JSON artifacts only.
"""
from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
import sqlite3
import statistics
import tempfile
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
GROUPS = [
    ("Vision", "tv", 11),
    ("Task Orientation", "to", 7),
    ("Support for Innovation", "si", 8),
    ("Participative Safety", "ps", 12),
]
COLUMNS = [f"{prefix}{i}" for _, prefix, size in GROUPS for i in range(1, size + 1)]


def write_csv(path, headers, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(headers)
        writer.writerows(rows)


def write_json(path, value):
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf-8")


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_jasp(path):
    with ZipFile(path) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        analyses = json.loads(archive.read("analyses.json"))["analyses"]
        if len(analyses) != 1:
            raise ValueError("Expected the single supplied CFA analysis.")
        analysis = analyses[0]
        with tempfile.TemporaryDirectory(prefix="tci-jasp-") as folder:
            db = Path(folder) / "internal.sqlite"
            db.write_bytes(archive.read("internal.sqlite"))
            connection = sqlite3.connect(db.as_uri() + "?mode=ro", uri=True)
            try:
                columns = connection.execute("SELECT id, name, columnType FROM Columns ORDER BY colIdx").fetchall()
                if [name for _, name, _ in columns] != COLUMNS:
                    raise ValueError("Unexpected JASP item order.")
                mappings = {
                    column_id: {
                        value: json.loads(original)
                        for value, original in connection.execute(
                            "SELECT value, originalValueJson FROM Labels WHERE columnId=?", (column_id,)
                        )
                    }
                    for column_id, _, _ in columns
                }
                # JASP stores category codes, not necessarily the displayed 1--5
                # response values. Decode each column using its own label map.
                query = "SELECT Filter_1," + ",".join(f"Column_{int(i)}" for i, _, _ in columns)
                query += " FROM DataSet_1 ORDER BY rowNumber"
                stored = connection.execute(query).fetchall()
                rows = []
                for record in stored:
                    if record[0] != 1:
                        raise ValueError("Unexpected filtered-out observation in the supplied project.")
                    rows.append([mappings[cid][value] for (cid, _, _), value in zip(columns, record[1:])])
            finally:
                connection.close()
    return manifest, analysis, rows, [kind for _, _, kind in columns]


def alpha(rows):
    k = len(rows[0])
    total_variance = statistics.variance([sum(row) for row in rows])
    return k / (k - 1) * (1 - sum(statistics.variance(c) for c in zip(*rows)) / total_variance)


def main():
    workbook_path = ROOT / "data/TCI.xlsx"
    jasp_path = ROOT / "analysis/CFA.jasp"
    book = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        sheet = book["TCI_148_B"]
        headers = [sheet.cell(1, c).value for c in range(1, 39)]
        codes = [sheet.cell(2, c).value for c in range(1, 39)]
        if codes != COLUMNS:
            raise ValueError("Unexpected workbook variable codes.")
        rows = [list(row) for row in sheet.iter_rows(min_row=3, max_row=150, max_col=38, values_only=True)]
        # The response matrix is A3:AL150. Do not use the shifted score formulas
        # or provenance labels elsewhere in the supplied workbook as row data.
        if len(rows) != 148 or any(type(v) is not int or v not in range(1, 6) for row in rows for v in row):
            raise ValueError("Expected 148 complete response rows with integer categories 1--5.")
        sample_formulas = {cell: sheet[cell].value for cell in ["AM2", "AM3", "AP2", "AQ2", "AR2"]}
        source_notes = [cell.value for cells in sheet.iter_rows(min_col=46, max_col=46, min_row=2) for cell in cells if cell.value]
    finally:
        book.close()

    manifest, analysis, jasp_rows, types = load_jasp(jasp_path)
    if rows != jasp_rows:
        raise ValueError("Workbook and JASP item responses differ (including row order).")

    write_csv(ROOT / "data/tci_items.csv", COLUMNS, rows)
    mapping = []
    offset = 0
    for dimension, _, size in GROUPS:
        for j in range(offset, offset + size):
            mapping.append([COLUMNS[j], f"TCI{j+1}", dimension, j+1, headers[j], 1, 5, False])
        offset += size
    write_csv(ROOT / "data/codebook.csv", ["variable", "manuscript_item", "dimension", "excel_column_number", "portuguese_item_header", "minimum", "maximum", "reverse_scored"], mapping)

    result_dir = ROOT / "results"
    result_dir.mkdir(exist_ok=True)
    frequencies = []
    for j, name in enumerate(COLUMNS):
        counts = [sum(row[j] == value for row in rows) for value in range(1, 6)]
        frequencies.append([name, f"TCI{j+1}", 148, *counts, *[count * 100 / 148 for count in counts]])
    write_csv(result_dir / "item_frequencies.csv", ["variable", "manuscript_item", "n", *[f"count_{v}" for v in range(1, 6)], *[f"percent_{v}" for v in range(1, 6)]], frequencies)

    score_rows = []
    dimension_scores = []
    offset = 0
    for dimension, _, size in GROUPS:
        subset = [row[offset:offset+size] for row in rows]
        scores = [statistics.mean(row) for row in subset]
        dimension_scores.append(scores)
        score_rows.append([dimension, len(scores), statistics.mean(scores), statistics.stdev(scores), statistics.median(scores), min(scores), max(scores), alpha(subset)])
        offset += size
    overall = [statistics.mean(values) for values in zip(*dimension_scores)]
    score_rows.append(["Overall (equal weight per dimension)", len(overall), statistics.mean(overall), statistics.stdev(overall), statistics.median(overall), min(overall), max(overall), ""])
    write_csv(result_dir / "score_descriptives.csv", ["score", "n", "mean", "sample_sd", "median", "minimum", "maximum", "raw_item_alpha"], score_rows)

    saved = analysis["results"]
    estimates = saved["estimates"]["collection"]
    fits = saved["maincontainer"]["collection"]
    tables = {
        "chi_square": fits["maincontainer_cfatab"]["data"],
        "fit_indices": fits["maincontainer_fits"]["collection"]["maincontainer_fits_indices"]["data"],
        "other_fit_measures": fits["maincontainer_fits"]["collection"]["maincontainer_fits_others"]["data"],
        "loadings": estimates["estimates_fl1"]["data"],
        "factor_correlations": estimates["estimates_fc"]["data"],
        "residual_variances": estimates["estimates_rv"]["data"],
        "residual_associations": estimates["estimates_rc"]["data"],
        "reliability": saved["resRelTable"]["data"],
    }
    write_json(result_dir / "jasp_saved_results.json", {
        "provenance": "Extracted saved JASP output; the CFA was not re-estimated by prepare_data.py.",
        "jasp_version": manifest["jaspVersion"],
        "analysis_module_version": analysis["dynamicModule"]["moduleVersion"],
        "options": analysis["options"],
        "tables": tables,
    })
    originals = [workbook_path, jasp_path, ROOT / "analysis/Script e syntax.docx"]
    write_json(result_dir / "validation.json", {
        "n": len(rows), "items": len(COLUMNS), "missing_item_values": 0,
        "response_values": sorted(set(v for row in rows for v in row)),
        "jasp_workbook_exact_match_including_order": True,
        "jasp_item_types": sorted(set(types)),
        "jasp_version": manifest["jaspVersion"],
        "analysis_module_version": analysis["dynamicModule"]["moduleVersion"],
        "cfa_status_in_saved_project": analysis["status"],
        "cfa_reestimated_by_this_script": False,
        "raw_38_item_alpha": alpha(rows),
        "workbook_formula_examples": sample_formulas,
        "workbook_notes_verbatim": source_notes,
        "saved_total_omega": tables["reliability"][-1]["rel"],
        "flags": [
            "The workbook identifies this reconstruction as Version B; the original selection record is not supplied.",
            "Workbook score formulas and source-row labels are shifted relative to the item-response rows; scores are recomputed from A3:AL150.",
            "Saved total omega exceeds one and must not be interpreted as valid reliability evidence.",
            "Original Python scripts for HTMT, descriptive comparisons and plotting were not included.",
            "No demographic variables or team identifiers are present in the shared item matrix.",
        ],
        "original_files_sha256": {str(p.relative_to(ROOT)).replace("\\", "/"): sha256(p) for p in originals},
    })
    print("Verified 148 complete rows x 38 ordinal items; workbook and JASP responses match exactly.")
    print("Exported item CSV, codebook, frequencies, descriptive summaries, and saved JASP tables.")
    print("See docs/DATA_NOTES.md for source-layout and analysis limitations.")


if __name__ == "__main__":
    main()
